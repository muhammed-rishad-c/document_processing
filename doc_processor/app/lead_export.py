import os
import json
import time
import threading
from pathlib import Path
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook

LEADS_DIR = Path(__file__).resolve().parent / "leads"
LEADS_DIR.mkdir(parents=True, exist_ok=True)

LEADS_OVERFLOW_LOG_PATH = Path(__file__).resolve().parent / "leads_overflow_log.jsonl"

_write_lock = threading.Lock()

COLUMN_HEADERS = ["Timestamp", "Question", "Name", "Email", "Phone", "Session ID"]

MAX_SAVE_ATTEMPTS = 3
RETRY_DELAY_SECONDS = 0.3


def _company_file_path(company_id: str) -> Path:
    return LEADS_DIR / f"{company_id}.xlsx"


def _create_new_workbook(company_name: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([f"Company: {company_name}"])
    ws.append(COLUMN_HEADERS)
    return wb


def _write_overflow(entry: dict, reason: str) -> None:
    """Last-resort log so a lead is never silently lost, even if the Excel
    write itself keeps failing (e.g. file open in Excel, disk issue)."""
    entry_with_reason = {**entry, "overflow_reason": reason}
    try:
        with open(LEADS_OVERFLOW_LOG_PATH, "a", encoding="utf-8") as f:
            f.write(json.dumps(entry_with_reason) + "\n")
    except Exception as e:
        # If even this fails, there's nothing more we can do locally —
        # surface it as loudly as possible.
        print(f"[lead_export] CRITICAL: failed to write overflow log too: {e}")


def append_lead(
    *,
    company_id: str,
    company_name: str,
    name: str,
    email: str,
    phone: str | None,
    question: str,
    session_id: str,
) -> None:
    """Append one lead row to leads/{company_id}.xlsx, creating the file
    with a header (including the company name) if it doesn't exist yet.
    Never raises — on repeated failure the row goes to a fallback overflow
    log instead, so a lead is never silently lost. Always logs a warning
    on any failure path."""

    row = [
        datetime.now(timezone.utc).isoformat(),
        question,
        name,
        email,
        phone or "",
        session_id,
    ]
    entry_for_overflow = {
        "timestamp": row[0],
        "company_id": company_id,
        "company_name": company_name,
        "question": question,
        "name": name,
        "email": email,
        "phone": phone,
        "session_id": session_id,
    }

    file_path = _company_file_path(company_id)

    last_error = None
    for attempt in range(1, MAX_SAVE_ATTEMPTS + 1):
        try:
            with _write_lock:
                if file_path.exists():
                    wb = load_workbook(file_path)
                    ws = wb["Leads"] if "Leads" in wb.sheetnames else wb.active
                else:
                    wb = _create_new_workbook(company_name)
                    ws = wb["Leads"]

                ws.append(row)
                wb.save(file_path)
            return  # success — done
        except Exception as e:
            last_error = e
            print(
                f"[lead_export] WARNING: attempt {attempt}/{MAX_SAVE_ATTEMPTS} "
                f"failed to save lead for company {company_id}: {e}"
            )
            if attempt < MAX_SAVE_ATTEMPTS:
                time.sleep(RETRY_DELAY_SECONDS)

    # All attempts failed — never lose the lead silently.
    print(
        f"[lead_export] WARNING: giving up on Excel write for company "
        f"{company_id} after {MAX_SAVE_ATTEMPTS} attempts, writing to "
        f"overflow log instead. Last error: {last_error}"
    )
    _write_overflow(entry_for_overflow, reason=str(last_error))
    
